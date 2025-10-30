"""Export metric data to XLSX using openpyxl."""
from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore

from ..logging_config import get_logger
from ..metrics.base import MetricResult

logger = get_logger(__name__)

def metrics_to_xlsx(results: Iterable[MetricResult], output_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Metrics"
    ws.append(["Key", "Label", "Value", "Description"])
    for r in results:
        ws.append([r.key, r.label, r.value, r.description or ""])
    # Autosize
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info("XLSX saved to %s", output_path)
    return output_path
